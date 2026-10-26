import os
import re
from pathlib import Path
from typing import Dict, Any, Tuple
import pypdf
import docx

class DocumentParser:
    """Handles multi-format document ingestion, text extraction, and normalization."""
    
    SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

    @classmethod
    def validate_file(cls, filename: str, file_bytes: bytes) -> Tuple[bool, str]:
        """Validate file extension and minimum byte size."""
        ext = Path(filename).suffix.lower()
        if ext not in cls.SUPPORTED_EXTENSIONS:
            return False, f"Unsupported file extension '{ext}'. Please upload a PDF, DOCX, or TXT file."
        
        if len(file_bytes) == 0:
            return False, "The uploaded file is empty (0 bytes). Please upload a valid document."
        
        if len(file_bytes) > 25 * 1024 * 1024: # 25 MB limit
            return False, "The file exceeds the 25 MB size limit."
            
        return True, ""

    @classmethod
    def extract_text(cls, filename: str, file_bytes: bytes) -> str:
        """Extract raw text from PDF, DOCX, or TXT file bytes."""
        ext = Path(filename).suffix.lower()
        text = ""

        if ext == ".txt" or ext == ".md":
            try:
                text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1", errors="replace")

        elif ext == ".pdf":
            import io
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            extracted_pages = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(page_text)
            text = "\n\n".join(extracted_pages)

        elif ext == ".docx":
            import io
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)

        return cls.normalize_text(text)

    @classmethod
    def normalize_text(cls, raw_text: str) -> str:
        """Clean and normalize extracted document text."""
        if not raw_text:
            return ""

        # Normalize line breaks
        text = raw_text.replace("\r\n", "\n").replace("\r", "\n")

        # De-hyphenate broken words across linebreaks e.g. "knowl-\nedge" -> "knowledge"
        text = re.sub(r'(\w+)-\n(\w+)', r'\1\2', text)

        # Replace non-breaking spaces and tabs
        text = text.replace('\xa0', ' ').replace('\t', ' ')

        # Remove strange non-printable control characters (keep standard punctuation and newlines)
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)

        # Collapse excessive whitespace
        text = re.sub(r' +', ' ', text)
        text = re.sub(r'\n{3,}', '\n\n', text)

        return text.strip()

    @classmethod
    def get_document_metrics(cls, text: str) -> Dict[str, Any]:
        """Compute basic structural metrics for the document."""
        words = re.findall(r'\b\w+\b', text)
        paragraphs = [p for p in text.split("\n\n") if p.strip()]
        # Approximate sentence count
        sentences = re.split(r'(?<=[.!?])\s+', text)
        sentences = [s.strip() for s in sentences if len(s.strip()) > 3]

        return {
            "char_count": len(text),
            "word_count": len(words),
            "sentence_count": len(sentences),
            "paragraph_count": len(paragraphs),
            "estimated_read_time_minutes": max(1, round(len(words) / 200))
        }
